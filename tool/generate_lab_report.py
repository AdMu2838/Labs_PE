"""Construye el informe DOCX desde la plantilla institucional y el código real."""
from __future__ import annotations

import argparse
import hashlib
import json
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / 'docs' / 'report_assets'
OUTPUT = ROOT / 'docs' / 'Informe_Laboratorio_4_UNSA_LAB.docx'
GARNET = '7B1835'


def diagram() -> None:
    image = Image.new('RGB', (1440, 740), '#F7F5F2')
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 24)
    bold = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 30)
    title = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 36)
    draw.text((55, 32), 'Arquitectura modular de UNSA · LAB', fill='#7B1835', font=title)
    nodes = [
        (55, 110, 660, 305, 'Presentación · Flutter Web', ['Pantallas y componentes Material 3', 'Riverpod: estado e inyección', 'GoRouter: navegación y protección de rutas']),
        (780, 110, 1385, 305, 'Dominio · Contratos y entidades', ['AuthRepository y LabRepository', 'AppSession, LabGroup y LabSession', 'Especificaciones de formularios']),
        (55, 410, 660, 625, 'Datos · Repositorios', ['SupabaseAuthRepository', 'SupabaseLabRepository', 'DemoLabRepository para mockups']),
        (780, 410, 1385, 625, 'Supabase · PostgreSQL', ['Auth: contraseña y Google OAuth', 'RLS, triggers y auditoría', 'RPC transaccionales, Realtime y catálogo']),
    ]
    for x1, y1, x2, y2, heading, lines in nodes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill='white', outline='#D8CEC7', width=2)
        draw.rounded_rectangle((x1, y1, x2, y1 + 62), radius=22, fill='#7B1835')
        draw.rectangle((x1, y1 + 30, x2, y1 + 62), fill='#7B1835')
        draw.text((x1 + 20, y1 + 15), heading, fill='white', font=bold)
        for i, line in enumerate(lines):
            draw.text((x1 + 24, y1 + 83 + i * 34), line, fill='#25242B', font=regular)
    for start, end in [((670, 205), (766, 205)), ((360, 315), (360, 395)), ((670, 520), (766, 520))]:
        draw.line((start, end), fill='#C5A15B', width=6)
        ex, ey = end
        if ex == start[0]:
            draw.polygon([(ex, ey), (ex - 12, ey - 18), (ex + 12, ey - 18)], fill='#C5A15B')
        else:
            draw.polygon([(ex, ey), (ex - 18, ey - 12), (ex - 18, ey + 12)], fill='#C5A15B')
    draw.text((55, 678), 'Las pantallas no ejecutan consultas SQL: utilizan contratos y repositorios.', fill='#6D6974', font=regular)
    image.save(ASSETS / '13_arquitectura.png')


def style(document, name, size, bold=False, color='25242B', outline=None):
    styles = document.styles
    item = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    item.base_style = styles['Normal']
    item.font.name = 'Aptos'
    item.font.size = Pt(size)
    item.font.bold = bold
    item.font.color.rgb = RGBColor.from_string(color)
    item.paragraph_format.space_after = Pt(6)
    item.paragraph_format.line_spacing = 1.08
    if bold:
        item.paragraph_format.keep_with_next = True
        item.paragraph_format.space_before = Pt(10)
    if outline is not None:
        element = OxmlElement('w:outlineLvl')
        element.set(qn('w:val'), str(outline))
        item.element.get_or_add_pPr().append(element)
    return item


def clear_cell(cell):
    for child in list(cell._tc):
        if child.tag != qn('w:tcPr'):
            cell._tc.remove(child)
    cell.add_paragraph()


def text(cell, value, style_name='Informe Texto', bold=False):
    if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text:
        p = cell.paragraphs[0]
        p.style = style_name
    else:
        p = cell.add_paragraph(style=style_name)
    p.paragraph_format.widow_control = True
    run = p.add_run(value)
    if bold:
        run.bold = True
    return p


def heading(cell, value, level=2):
    p = text(cell, value, f'Informe H{level}')
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    return p


def comment(document, paragraph, value):
    document.add_comment(paragraph.runs, text=value, author='Equipo UNSA · LAB', initials='AU/JG')


def hyperlink(paragraph, label, url):
    relationship = paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    element = OxmlElement('w:hyperlink')
    element.set(qn('r:id'), relationship)
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    content = OxmlElement('w:t')
    content.text = label
    run.append(content)
    element.append(run)
    paragraph._p.append(element)


def figure(cell, number, filename, caption, description):
    p = cell.add_paragraph(style='Informe Figura')
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / filename), width=Cm(16.2))
    cap = text(cell, f'Figura {number}. {caption}', 'Informe Leyenda')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    text(cell, description)


def set_existing(cell, value, bold=False):
    sample_p = cell.paragraphs[0]
    ppr = deepcopy(sample_p._p.pPr) if sample_p._p.pPr is not None else None
    sample_run = next((r for p in cell.paragraphs for r in p.runs if r.text), None)
    rpr = deepcopy(sample_run._r.rPr) if sample_run is not None and sample_run._r.rPr is not None else None
    clear_cell(cell)
    p = cell.paragraphs[0]
    if ppr is not None:
        p._p.insert(0, ppr)
    run = p.add_run(value)
    if rpr is not None:
        run._r.insert(0, rpr)
    run.font.name = 'Aptos'
    run.font.size = Pt(12)
    run.bold = bold


def appendix_files():
    groups = [
        ('A. Código completo de Flutter', sorted((ROOT / 'lib').rglob('*.dart'))),
        ('B. Base de datos y reglas de negocio', sorted((ROOT / 'supabase' / 'migrations').glob('*.sql'))),
        ('C. Pruebas de la aplicación', sorted((ROOT / 'test').glob('*.dart'))),
        ('D. Configuración y ejecución', [ROOT / p for p in [
            'pubspec.yaml', 'analysis_options.yaml', 'config.example.json',
            'web/index.html', 'web/manifest.json', 'web/favicon.svg', 'web/_redirects',
            'firebase.json', '.gitignore', 'tool/serve_web.py', 'tool/check_supabase.mjs',
        ]]),
    ]
    return groups


def paginate_results(document, cell):
    """Expande el bloque de resultados en filas del mismo formato, sin cortar figuras."""
    row = cell._tc.getparent()
    parent = row.getparent()
    headings = {document.styles[f'Informe H{i}'].style_id for i in [1, 2, 3]}
    groups = []
    current = []
    has_body = False
    for p in list(cell._tc):
        if p.tag != qn('w:p'):
            continue
        styles = p.xpath('./w:pPr/w:pStyle')
        is_heading = bool(styles and styles[0].get(qn('w:val')) in headings)
        if is_heading and current and has_body:
            groups.append(current)
            current = []
            has_body = False
        for page_break in p.xpath('./w:pPr/w:pageBreakBefore'):
            page_break.getparent().remove(page_break)
        current.append(p)
        has_body = has_body or not is_heading
    if current:
        groups.append(current)
    template_row = deepcopy(row)
    for group in groups:
        next_row = deepcopy(template_row)
        properties = next_row.get_or_add_trPr()
        for element in list(properties):
            if element.tag in [qn('w:trHeight'), qn('w:cantSplit')]:
                properties.remove(element)
        contains_image = any(p.xpath('.//w:drawing') for p in group)
        length = sum(len(''.join(p.xpath('.//w:t/text()'))) for p in group)
        if contains_image or length < 1500:
            properties.append(OxmlElement('w:cantSplit'))
        next_cell = next_row.tc_lst[0]
        for element in list(next_cell):
            if element.tag != qn('w:tcPr'):
                next_cell.remove(element)
        for p in group:
            next_cell.append(p)
        parent.insert(parent.index(row), next_row)
    parent.remove(row)


def build(template):
    if not template.exists():
        raise SystemExit(f'No se encuentra la plantilla: {template}')
    ASSETS.mkdir(parents=True, exist_ok=True)
    required = [f'{i:02d}_{name}_web.png' for i, name in enumerate([
        'login', 'dashboard', 'catalogo', 'confirmacion', 'matriculas', 'horario',
        'plan', 'solicitud', 'admin', 'formulario_grupo', 'importacion', 'reportes',
    ], 1)]
    for filename in required:
        with Image.open(ASSETS / filename) as image:
            if image.size != (1440, 900):
                raise ValueError(f'La captura no corresponde a escritorio: {filename} {image.size}')
    diagram()
    document = Document(template)
    style(document, 'Informe Texto', 12)
    style(document, 'Informe H1', 12, True, '000000', 0)
    style(document, 'Informe H2', 12, True, '000000', 1)
    style(document, 'Informe H3', 11, True, '000000', 2)
    style(document, 'Informe Leyenda', 10, False, '5A5350')
    style(document, 'Informe Figura', 10)
    code_style = style(document, 'Informe Código', 8.5, False, '000000')
    code_style.font.name = 'Courier New'
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = Pt(10)
    code_style.paragraph_format.keep_with_next = False
    code_style.paragraph_format.widow_control = False
    code_style.paragraph_format.tab_stops.add_tab_stop(Cm(.6))
    table = document.tables[0]
    # Se mantienen la tabla, su cuadrícula, bordes y secciones de la plantilla.
    set_existing(table.rows[2].cells[1], 'Sistema web de matrícula de laboratorios con Flutter y Supabase — UNSA · LAB')
    set_existing(table.rows[3].cells[1], '4')
    set_existing(table.rows[3].cells[3], '2026')
    set_existing(table.rows[3].cells[6], 'X')
    set_existing(table.rows[4].cells[1], 'Por completar')
    set_existing(table.rows[4].cells[3], 'Por completar')
    members = table.rows[5].cells[0]
    clear_cell(members)
    text(members, 'INTEGRANTE(s)', bold=True)
    text(members, 'Mamani Uscamayta Agustin David\nCusilayme Garcia Jose Luis')
    set_existing(table.rows[5].cells[6], 'Por evaluar')
    for index in [9, 10, 13, 16, 19]:
        clear_cell(table.rows[index].cells[0])
        for element in table.rows[index]._tr.xpath('./w:trPr/w:cantSplit'):
            element.getparent().remove(element)
    for index in [8, 12, 15, 18]:
        for p in table.rows[index].cells[0].paragraphs:
            p.paragraph_format.keep_with_next = True

    results = table.rows[9].cells[0]
    heading(results, 'EJERCICIOS Y PROBLEMAS PROPUESTOS RESUELTOS', 1)
    text(results, 'Enunciado: desarrollar una aplicación web para matricular estudiantes en los laboratorios de la Escuela Profesional de Ingeniería de Sistemas de la UNSA, considerando dificultades habituales del proceso y funciones para estudiantes y administradores. Se implementó una solución con Flutter Web y Supabase, ampliada con roles de docente y coordinación.')
    p = text(results, 'Repositorio del proyecto: ')
    hyperlink(p, 'https://github.com/AdMu2838/Labs_PE', 'https://github.com/AdMu2838/Labs_PE')
    text(results, 'La carpeta local se inicializó en Git y el remoto origin quedó configurado hacia este repositorio. El documento incorpora el código fuente completo para permitir su revisión junto con las evidencias de la aplicación.')

    heading(results, '1. Descripción general del programa')
    text(results, 'UNSA · LAB centraliza la consulta de grupos, vacantes, horarios y requisitos académicos para la matrícula de laboratorios. Su finalidad es reducir la búsqueda manual de horarios, evitar grupos duplicados y cruces, y proporcionar trazabilidad de cambios y retiros. La interfaz se presenta en español y utiliza una paleta granate, dorado y fondos claros inspirada en la identidad de la universidad.')
    text(results, 'Flutter desarrolla la interfaz y la navegación; Supabase proporciona autenticación, PostgreSQL y actualización de información mediante Realtime. La malla curricular se obtuvo del plan oficial 2025 de Ingeniería de Sistemas: 87 asignaturas y 72 relaciones de prerrequisito. Se registran códigos, semestres, créditos y horas de laboratorio. Los 302 créditos codificados incluyen todas las alternativas electivas y no equivalen a la carga de un estudiante [1], [5], [8].')
    heading(results, '1.1. Funcionalidades por rol', 3)
    for value in [
        'Estudiante: registro con CUI, login por contraseña o Google, recuperación de contraseña, consulta y filtrado de grupos, matrícula, cambio de grupo, retiro, lista de espera, horario, notificaciones, solicitudes y constancia PDF.',
        'Administrador: administración de usuarios y roles, periodos, ambientes, grupos, horarios, planes, asignaturas, prerrequisitos, habilitaciones e historial; importación CSV, reportes y auditoría.',
        'Coordinación: gestión de la oferta y datos académicos, habilitación de cursos, importación y atención de solicitudes. La asignación de roles corresponde al administrador.',
        'Docente: consulta de sus grupos y estudiantes matriculados, horario, exportación de listas y publicación de avisos a grupos asignados.',
    ]:
        text(results, value)
    heading(results, '1.2. Reglas principales', 3)
    text(results, 'La base de datos verifica que exista una cuenta activa, CUI, plan compatible, periodo abierto, horario publicado y habilitación cuando corresponda. Impide una segunda matrícula activa del mismo curso en un periodo, valida prerrequisitos y créditos mínimos, y compara todas las sesiones para detectar cruces. Un bloqueo transaccional por periodo impide que dos solicitudes utilicen la última vacante. El cambio de grupo se realiza de manera atómica: si el destino no puede aceptarse, se conserva la matrícula de origen.')
    text(results, 'La cola de espera se ordena por fecha de registro y se revisa al liberar o ampliar vacantes dentro de la ventana de matrícula. Cada acción relevante genera notificaciones y eventos de auditoría. Los permisos se aplican mediante Row Level Security (RLS) y funciones PostgreSQL; ocultar una opción de la interfaz no es la única medida de autorización [6].')

    heading(results, '2. Diseño de interfaces: funcionalidades y mockups')
    text(results, 'El mockup se construyó como una interfaz funcional y modular. La navegación lateral permanece visible en escritorio y la zona central presenta tarjetas, filtros y formularios. El sistema visual se centraliza en app_theme.dart, con granate #7B1835, granate oscuro #4C1026, dorado #C5A15B y fondo #F7F5F2. Estos códigos son una interpretación visual de la identidad universitaria, no una certificación de un manual de marca.')
    text(results, 'Todas las capturas de este informe corresponden a vista web de escritorio de 1440 × 900 píxeles. La figura 1 muestra el login conectado al proyecto Supabase. Las figuras 2–12 usan el modo demostración de la propia aplicación para ilustrar el diseño con datos ficticios reproducibles; la franja visible identifica ese modo. Las verificaciones remotas reales se describen en la sección de pruebas, y no se atribuyen a los datos de los mockups.')

    figures = [
        (1, '01_login_web.png', 'Login web con contraseña y acceso Google.', 'La pantalla separa la identificación visual del formulario. TextFormField permite capturar correo y contraseña; Form agrupa la validación, el botón de visibilidad cambia obscureText y los botones inician registro, recuperación o autenticación OAuth. El panel de acceso limita su ancho mediante ConstrainedBox. La autenticación es administrada por Supabase Auth, sin solicitar la contraseña de Google a la aplicación.'),
        (2, '02_dashboard_web.png', 'Panel principal y resumen de matrícula.', 'El dashboard presenta un saludo, acceso al catálogo, indicadores y calendario. Container aplica el degradado institucional; Row y Column organizan el contenido; Wrap distribuye las tarjetas según el ancho. Panel reutiliza Card y Padding para mantener un mismo estilo. Los proveedores de Riverpod permiten refrescar las cifras y mostrar carga o error cuando se consultan datos.'),
        (3, '03_catalogo_web.png', 'Catálogo de grupos con filtros y vacantes.', 'TextField filtra por curso, código o docente. DropdownButtonFormField selecciona semestre y periodo, mientras FilterChip activa la condición de vacantes disponibles. LayoutBuilder obtiene el ancho y Wrap distribuye las tarjetas. StatusChip distingue estados; LinearProgressIndicator muestra la ocupación. FilledButton comunica la acción de matrícula o ingreso a lista de espera.'),
        (4, '04_confirmacion_web.png', 'Confirmación de matrícula en escritorio.', 'showDialog presenta un AlertDialog con el curso, grupo, ambiente y sesiones. La función confirm devuelve una decisión booleana antes de invocar la operación. La interfaz deshabilita envíos mientras se procesa una solicitud y SnackBar comunica el resultado. La validación definitiva ocurre en la función enroll_lab de PostgreSQL.'),
        (5, '05_matriculas_web.png', 'Matrículas, cambio de grupo y constancia.', 'La vista muestra registros activos y retirados con su periodo, ambiente y horario. OutlinedButton inicia un SimpleDialog para elegir un destino; TextButton solicita el retiro. La exportación utiliza DocumentService para generar un PDF de consulta. RecordsView gestiona la lectura asíncrona y EmptyState explica los casos sin resultados.'),
        (6, '06_horario_web.png', 'Horario semanal de laboratorios.', 'El horario agrupa sesiones por día. LayoutBuilder calcula el número de columnas y Wrap distribuye los paneles. Cada sesión utiliza Container con borde lateral para destacar curso, hora y ambiente. En escritorio pueden mostrarse cinco días en paralelo. La comparación de intervalos de LabSession también permite interpretar por qué dos sesiones consecutivas no constituyen un cruce.'),
        (7, '07_plan_web.png', 'Consulta del plan de estudios por semestre.', 'ExpansionTile permite expandir cada semestre sin saturar la pantalla. ListTile presenta nombre, código, créditos y horas de laboratorio. TextField busca asignaturas y SwitchListTile restringe los resultados a cursos con laboratorio. La base real carga el catálogo oficial completo; el repositorio de demostración contiene una selección de cursos para el mockup.'),
        (8, '08_solicitud_web.png', 'Formulario de solicitudes del estudiante.', 'RecordForm compone un Form a partir de FormFieldSpec. TextFormField recoge asunto y descripción; DropdownButtonFormField selecciona el tipo. SingleChildScrollView evita recortar el formulario y los botones Guardar/Cancelar controlan la operación. Las solicitudes quedan pendientes hasta que coordinación registra una respuesta; su aprobación no omite las reglas de matrícula.'),
        (9, '09_admin_web.png', 'Panel de administración modular.', 'ChoiceChip selecciona periodos, ambientes, grupos y otros mantenimientos. La definición de AdminSection evita duplicar el diseño de cada CRUD. RecordsView carga registros y Panel agrupa las acciones de edición y eliminación. GoRouter comprueba el rol para acceder a la ruta, mientras RLS aplica los permisos en el servidor.'),
        (10, '10_formulario_grupo_web.png', 'Creación de grupos y asignación de recursos.', 'Los desplegables consultan repositorios para seleccionar asignatura, periodo, laboratorio y docente. TextFormField valida vacantes y SwitchListTile controla la publicación. El formulario se muestra como AlertDialog en escritorio. PostgreSQL comprueba capacidad física, asignatura con laboratorio y conflictos de horario o docente antes de guardar.'),
        (11, '11_importacion_web.png', 'Importación CSV de datos académicos.', 'ImportPanel utiliza FilePicker para seleccionar un CSV, muestra una previsualización y verifica cabeceras y número de registros. El tipo de datos se selecciona con DropdownButtonFormField. La función import_academic resuelve CUI, código y periodo y procesa el lote en una transacción: una fila inválida cancela el conjunto.'),
        (12, '12_reportes_web.png', 'Reportes de ocupación y exportación.', 'Cada grupo muestra matriculados, capacidad y vacantes mediante una barra de progreso. OutlinedButton activa la descarga CSV. DocumentService separa la generación del archivo de la pantalla. Los datos individuales exportables se limitan a los registros autorizados para el docente o administrador.'),
    ]
    for number, filename, caption, description in figures:
        heading(results, f'2.{number}. {caption}', 3)
        figure(results, number, filename, caption, description)

    heading(results, '3. Descripción del proceso de desarrollo')
    heading(results, '3.1. Análisis del problema y diseño académico', 3)
    text(results, 'Primero se identificaron problemas de matrícula: información dispersa, ausencia de vacantes, cruces de horario, repetición de inscripciones y cambios sin trazabilidad. Se definieron roles y casos de uso y se revisó la fuente oficial de la UNSA. Después se separaron la malla curricular permanente y la oferta por periodo: las asignaturas no se confunden con grupos, ambientes o sesiones de un semestre particular.')
    heading(results, '3.2. Arquitectura y organización del código', 3)
    text(results, 'El código se organiza por funcionalidades y aplica separación de responsabilidades inspirada en una arquitectura por capas. app contiene configuración, rutas, shell y tema; core define contratos, configuración e implementaciones compartidas; shared reúne componentes reutilizables; features agrupa autenticación, matrícula, administración, horarios, perfiles y reportes. El módulo auth contiene domain, data y presentation. Otros módulos reutilizan el contrato transversal LabRepository y las entidades de matrícula en lugar de ejecutar consultas en sus pantallas.')
    figure(results, 13, '13_arquitectura.png', 'Separación de responsabilidades de la aplicación.', 'La presentación depende de contratos e inyección de dependencias. SupabaseLabRepository ejecuta consultas y RPC; DemoLabRepository implementa el mismo contrato con datos temporales. Este diseño permite construir y verificar el mockup sin credenciales, manteniendo la integración real en una implementación separada.')
    heading(results, '3.3. Construcción de widgets reutilizables', 3)
    for title, body in [
        ('Scaffold, Material y AppBar', 'Scaffold organiza las zonas principales; Material integra superficies e interacciones Material 3. AppBar y el shell muestran acceso a cuenta, actualización y navegación. En escritorio se agrega un menú lateral; con menor ancho se utiliza Drawer.'),
        ('Row, Column y Expanded', 'Distribuyen los elementos en ejes horizontal y vertical. Expanded asigna el espacio restante y limita el texto en una fila para evitar desbordes. No se fijan alturas rígidas para los textos del contenido.'),
        ('LayoutBuilder, Wrap y ConstrainedBox', 'LayoutBuilder trabaja con las restricciones reales del contenedor; Wrap cambia de línea cuando falta ancho; ConstrainedBox limita el contenido máximo. El catálogo usa una, dos o tres columnas según espacio disponible [1], [2].'),
        ('SingleChildScrollView y PageBody', 'Cada ruta posee su propio scroll dentro de PageBody. El Navigator mantiene el tamaño del viewport; ubicar el scroll en cada página permite navegar entre módulos de diferente altura sin desbordes durante las transiciones.'),
        ('Form y campos de entrada', 'Form utiliza GlobalKey<FormState> y los validadores de TextFormField para comprobar datos. TextEditingController conserva la entrada y debe liberarse en dispose. Los formularios incluyen controles para correo, contraseña, CUI, texto, números, fechas y horas.'),
        ('Controles de selección', 'DropdownButtonFormField selecciona valores o recursos; ChoiceChip cambia de mantenimiento; FilterChip representa un filtro; SwitchListTile gestiona valores booleanos. RecordForm deriva estos controles de una especificación declarativa.'),
        ('Card, Container, Padding y Panel', 'Card aporta la superficie, Container configura bordes y colores y Padding controla el espacio interno. Panel encapsula estos elementos y facilita la uniformidad visual entre módulos.'),
        ('Diálogos, indicadores y mensajes', 'AlertDialog y SimpleDialog solicitan confirmaciones o selecciones. CircularProgressIndicator comunica carga, LinearProgressIndicator expresa ocupación y SnackBar devuelve el resultado. EmptyState evita presentar espacios vacíos sin explicación.'),
        ('ConsumerWidget y ConsumerStatefulWidget', 'Acceden a proveedores con ref.watch para reconstruir la interfaz cuando cambia la información. Los FutureProvider representan consultas; los NotifierProvider administran sesión demo y revisión; los repositorios se inyectan mediante Provider [3].'),
        ('GoRouter y ShellRoute', 'GoRouter define rutas URL y redirecciones según sesión, rol y CUI. ShellRoute conserva el marco común del campus. La estrategia de URL por ruta facilita el retorno OAuth y exige fallback a index.html en el hosting [4].'),
    ]:
        text(results, f'{title}. {body}')
    heading(results, '3.4. Base de datos y transacciones', 3)
    text(results, 'Se implementaron 16 tablas: plans, profiles, courses, prerequisites, periods, rooms, groups, sessions, academic_records, eligibilities, enrollments, waitlist, requests, notifications, announcements y audit_logs. Los UUID relacionan perfiles con Auth y conectan asignaturas, periodos y grupos. Se utilizan claves foráneas, checks e índices únicos parciales para impedir inconsistencias como una matrícula activa duplicada.')
    text(results, 'Las migraciones se aplicaron al proyecto MatriculasLab y quedaron registradas en supabase_migrations.schema_migrations. La migración 001 crea estructura, políticas y triggers; 002 agrega las operaciones y Realtime; 003 carga el plan 2025; 004 incorpora el alta sin CUI y la asignación inicial protegida. Las funciones enroll_lab, withdraw_lab y promote_queue ejecutan las operaciones críticas dentro de PostgreSQL.')
    text(results, 'En una matrícula se autentica al estudiante, se bloquea el periodo, se valida su situación y se cuenta la ocupación. Si existen vacantes se inserta la matrícula; si el grupo está lleno y se solicita espera, se registra en la cola. Un cambio identifica la matrícula anterior y comprueba el destino antes de sustituirla. Ante un error se revierte la operación. La prueba con dos solicitudes simultáneas a una única vacante confirmó que solo una obtuvo matrícula.')
    heading(results, '3.5. Autenticación y Google OAuth', 3)
    p = text(results, 'El login por correo utiliza Supabase Auth. El acceso Google se agregó mediante signInWithOAuth(OAuthProvider.google), con PKCE y retorno al origen de Flutter. El proveedor se habilitó en Supabase y se comprobó una redirección 302 hacia accounts.google.com con el callback del proyecto. Los usuarios nuevos reciben rol estudiante y, si no tienen CUI, son enviados a Mi perfil. PostgreSQL permite asignarlo una vez y bloquea cambios posteriores no autorizados [7].')
    comment(document, p, 'Cambio realizado: se añadió el botón Continuar con Google, comprobación del proveedor y retorno OAuth. También se incorporó la migración 004 para completar el CUI de una cuenta nueva. La prueba de consentimiento personal completo quedó pendiente por decisión del usuario.')
    text(results, 'Las variables SUPABASE_URL y SUPABASE_PUBLISHABLE_KEY se cargan mediante dart-define-from-file. La clave pública identifica el proyecto; la sesión del usuario y las políticas RLS determinan el acceso. El Client Secret de Google permanece en Supabase y no se agrega a la aplicación web.')
    heading(results, '3.6. Comprobaciones y resultados', 3)
    text(results, 'Se aprobaron 12 pruebas locales de widgets, accesibilidad y lógica, y flutter analyze no reportó incidencias. La compilación web release con la configuración del nuevo proyecto finalizó correctamente. En la validación inicial de Supabase real se completaron 44 comprobaciones positivas: autenticación, RLS, CRUD administrativo, concurrencia, prerrequisitos, cambios atómicos, cola, solicitudes, notificaciones, auditoría y registro único del CUI. También se verificaron desde Flutter la carga de datos remotos y descargas CSV y PDF.')
    text(results, 'La validación inicial conservó 87 asignaturas, 72 prerrequisitos, 302 créditos codificados, cuatro migraciones y 16 tablas con RLS. Los cinco usuarios temporales y sus registros de esa primera etapa se eliminaron. Posteriormente se crearon tres cuentas persistentes de estudiante, docente y administrador, y se ejecutaron 45 comprobaciones end-to-end con sesiones independientes de navegador. La oferta LAB04 - PRUEBAS contiene cinco grupos y nueve sesiones; el estudiante conserva tres matrículas compatibles. El informe docs/e2e_roles.md documenta esta nueva etapa, cuyas cuentas permanecen disponibles.')
    text(results, 'Alcance de la verificación: se comprobó la apertura de la pantalla oficial de Google y su callback, pero no se completó el consentimiento con una cuenta personal. La entrega de confirmaciones y recuperación por correo no se verificó: no hay SMTP personalizado y Supabase rechazó las direcciones temporales usadas. La integración académica utiliza importación administrativa, no una conexión automática con el sistema académico UNSA.')
    heading(results, '3.7. Registro comentado de modificaciones', 3)
    changes = [
        ('Migración a la cuenta Supabase indicada', 'Se sustituyó la configuración local del proyecto y se ejecutaron las cuatro migraciones en MatriculasLab. Se comprobó el esquema antes de aplicar las operaciones y se registró el historial de versiones.'),
        ('Incorporación de Google Auth y CUI', 'Se agregó acceso Google, comprobación del proveedor, redirección al perfil cuando falta CUI y restricción de asignación única en PostgreSQL.'),
        ('Ajuste del contenedor de páginas', 'Se trasladó el desplazamiento a PageBody para que cada ruta mantenga su propio scroll y evitar desbordes al navegar entre páginas de diferente altura.'),
        ('Preparación de evidencias e informe', 'Se capturaron únicamente vistas de escritorio, se documentaron los resultados reales y se anexó el código íntegro. Los cambios se registran como comentarios de Word junto a esta sección.'),
        ('Accesibilidad y validación end-to-end', 'Se separaron los nodos accesibles de Panel y Card para identificar los controles de forma independiente. Se agregó una prueba de accesibilidad y se comprobaron 45 casos con las tres cuentas persistentes, incluida matrícula, espera, importación y descargas.'),
    ]
    for title, body in changes:
        p = text(results, f'{title}: {body}')
        comment(document, p, body)

    questions = table.rows[10].cells[0]
    heading(questions, 'CUESTIONARIO TÉCNICO DEL DESARROLLO', 1)
    for question, answer in [
        ('¿Cómo se logra que el mockup se adapte al espacio disponible?', 'Se evalúan las restricciones con LayoutBuilder y se distribuyen elementos con Wrap y Expanded. Cada página permite scroll; el shell cambia la navegación lateral por Drawer cuando disminuye el ancho. La verificación local incluyó tamaños de 360, 768 y 1440 px y texto ampliado, aunque las evidencias gráficas del informe son únicamente web de escritorio.'),
        ('¿Por qué las validaciones no se realizan solo en Flutter?', 'Un cliente puede modificar solicitudes o ejecutar operaciones directamente. Las funciones PostgreSQL y RLS aplican permisos, requisitos y vacantes independientemente del diseño visual, y una transacción garantiza consistencia ante peticiones simultáneas.'),
        ('¿Qué ventaja aporta utilizar repositorios?', 'Un contrato uniforme permite reemplazar la fuente de información real por datos de demostración sin reescribir las pantallas. También concentra consultas y operaciones de datos fuera de los widgets.'),
        ('¿Cómo se interpreta la prueba de Google y de los correos?', 'La redirección y el callback Google quedaron configurados y comprobados. No se afirma que se haya completado el consentimiento de una cuenta personal. Las operaciones por correo necesitan verificar destinatarios y SMTP antes de considerarse una prueba de entrega exitosa.'),
    ]:
        heading(questions, question, 3)
        text(questions, answer)

    conclusions = table.rows[13].cells[0]
    heading(conclusions, '4. Apreciación de la herramienta')
    heading(conclusions, 'Dos características favorables de Flutter', 3)
    text(conclusions, '1. Composición y reutilización de interfaces. La jerarquía de widgets permite construir Panel, RecordForm y PageBody reutilizables. Un mismo diseño conserva consistencia en login, catálogo y administración y facilita la adaptación al tamaño disponible.')
    text(conclusions, '2. Productividad durante el desarrollo. Hot reload permite observar cambios del mockup rápidamente; las herramientas de análisis y pruebas de widgets ayudan a detectar errores de navegación y desbordes antes de entregar la compilación web.')
    heading(conclusions, 'Dos características desfavorables de Flutter', 3)
    text(conclusions, '1. Costo inicial de carga y distribución web. El navegador debe cargar el motor y los recursos de la aplicación. Frente a una página HTML simple, el arranque y el tamaño de distribución pueden ser mayores; se debe revisar la compilación release y limitar recursos innecesarios.')
    text(conclusions, '2. Complejidad de restricciones, estado y semántica web. Las filas y columnas mal restringidas producen desbordes y una pantalla dinámica exige comprender el ciclo de vida y el estado. Además, la semántica de widgets renderizados no siempre coincide con un DOM HTML tradicional, lo que exige cuidados de accesibilidad y automatización.')
    heading(conclusions, 'Conclusiones del trabajo', 3)
    text(conclusions, 'El resultado es una plataforma modular que integra la consulta y matrícula de laboratorios con una base académica verificable. La separación entre interfaz y reglas transaccionales permitió representar el proceso en un mockup funcional y comprobarlo contra Supabase real. Las pruebas de concurrencia y cambio de grupo respaldan el tratamiento de los principales problemas de matrícula.')
    text(conclusions, 'El diseño reutilizable permitió conservar la identidad visual y ampliar módulos sin concentrarlos en una sola pantalla. Como trabajo posterior se debe completar la prueba personal de Google, configurar correo para los estudiantes y registrar la oferta académica real de cada periodo. Estas tareas se diferencian de las capacidades implementadas y probadas.')

    method = table.rows[16].cells[0]
    for value in [
        '1. Analizar problemas del proceso de matrícula y definir usuarios, reglas y alcance.',
        '2. Consultar el plan de estudios oficial y modelar asignaturas, prerrequisitos y oferta por periodo.',
        '3. Preparar Flutter Web, organizar módulos y diseñar el tema institucional.',
        '4. Construir navegación, componentes y mockups con repositorios de demostración.',
        '5. Implementar Supabase Auth, RLS, triggers y operaciones transaccionales.',
        '6. Aplicar migraciones, importar el plan y comprobar permisos y concurrencia con datos temporales.',
        '7. Incorporar Google OAuth y el registro inicial de CUI, manteniendo roles protegidos.',
        '8. Ejecutar análisis, pruebas y compilación; revisar exportación y navegación desde el navegador.',
        '9. Vincular Git, capturar vistas web y elaborar el informe con comentarios y código completo.',
    ]:
        text(method, value)

    references = table.rows[19].cells[0]
    citations = [
        ('Flutter', 'Flutter documentation', 'https://docs.flutter.dev/'),
        ('Flutter', 'LayoutBuilder y API de widgets', 'https://api.flutter.dev/flutter/widgets/LayoutBuilder-class.html'),
        ('Riverpod', 'Documentación de Riverpod', 'https://riverpod.dev/docs/introduction/getting_started'),
        ('Flutter Packages', 'go_router', 'https://pub.dev/packages/go_router'),
        ('Supabase', 'Build a User Management App with Flutter', 'https://supabase.com/docs/guides/getting-started/tutorials/with-flutter'),
        ('Supabase', 'Row Level Security', 'https://supabase.com/docs/guides/database/postgres/row-level-security'),
        ('Supabase', 'Login with Google', 'https://supabase.com/docs/guides/auth/social-login/auth-google'),
        ('UNSA, Escuela Profesional de Ingeniería de Sistemas', 'Plan de estudios 2025', 'https://fips.unsa.edu.pe/ingenieriadesistemas/plan-de-estudios-2025/'),
        ('A. D. Mamani Uscamayta y J. L. Cusilayme Garcia', 'Repositorio Labs_PE', 'https://github.com/AdMu2838/Labs_PE'),
    ]
    for i, (author, title, url) in enumerate(citations, 1):
        p = text(references, f'[{i}] {author}, «{title}». [En línea]. Disponible: ')
        hyperlink(p, url, url)
        p.add_run('. [Consulta: 11 de octubre de 2026].')

    paginate_results(document, results)

    # Los anexos conservan la sección A4, el encabezado, las imágenes y el pie originales.
    p = document.add_paragraph(style='Informe H1')
    p.paragraph_format.page_break_before = True
    p.add_run('5. ANEXOS — CÓDIGO DE LA APLICACIÓN')
    p = document.add_paragraph(style='Informe Texto')
    p.add_run('Se incorpora el contenido íntegro de los archivos de implementación, migraciones, pruebas y configuración. Las líneas largas se ajustan visualmente en Word sin eliminar su contenido. Los archivos generados por Flutter y las credenciales locales no forman parte del código reproducible de este anexo.')
    source_manifest = []
    file_count = 0
    line_count = 0
    for group_title, files in appendix_files():
        document.add_paragraph(group_title, style='Informe H2')
        for file in files:
            relative = file.relative_to(ROOT).as_posix()
            source = file.read_text(encoding='utf-8-sig')
            file_count += 1
            line_count += len(source.splitlines())
            head = document.add_paragraph(f'{file_count:02d}. {relative}', style='Informe H3')
            head.paragraph_format.keep_with_next = True
            code_table = document.add_table(rows=1, cols=1)
            code_table.style = table.style
            code_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            code_table.autofit = False
            code_table.columns[0].width = Cm(17.1)
            cell = code_table.cell(0, 0)
            cell.width = Cm(17.1)
            for index, line in enumerate(source.splitlines()):
                p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
                p.style = 'Informe Código'
                p.add_run(line if line else ' ')
            spacer = document.add_paragraph()
            spacer.paragraph_format.space_after = Pt(3)
            spacer.paragraph_format.space_before = Pt(0)
            source_manifest.append({'file': relative, 'lines': len(source.splitlines()), 'sha256': hashlib.sha256(source.encode('utf-8')).hexdigest()})
    settings = document.settings.element
    update = settings.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields')
        settings.append(update)
    update.set(qn('w:val'), 'true')
    document.core_properties.title = 'Informe de laboratorio 4 — UNSA · LAB'
    document.core_properties.author = 'Mamani Uscamayta Agustin David; Cusilayme Garcia Jose Luis'
    document.core_properties.subject = 'Plataformas Emergentes — Semestre X'
    document.core_properties.comments = 'Informe con formato institucional, vistas web, comentarios de cambios y código completo.'
    document.save(OUTPUT)
    manifest = {'output': OUTPUT.name, 'figures': 13, 'web_screenshots': 12, 'source_files': file_count, 'source_lines': line_count, 'sources': source_manifest}
    (ASSETS / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in manifest.items() if k != 'sources'}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--template', type=Path, required=True)
    args = parser.parse_args()
    build(args.template)
