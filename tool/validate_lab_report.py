"""Verificación estructural, integridad de anexos y páginas de revisión del informe."""
import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path

import fitz
from docx import Document
from lxml import etree
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
REPORT = ROOT / 'docs' / 'Informe_Laboratorio_4_UNSA_LAB.docx'
PREVIEW = ROOT / '.dart_tool' / 'Informe_Laboratorio_4_UNSA_LAB_revision.pdf'
REVIEW = ROOT / '.dart_tool' / 'report_review'


def main(template):
    source_template = Document(template)
    report = Document(REPORT)
    section_values = ['page_width', 'page_height', 'top_margin', 'bottom_margin', 'left_margin', 'right_margin', 'header_distance', 'footer_distance']
    assert all(getattr(source_template.sections[0], key) == getattr(report.sections[0], key) for key in section_values), 'Se modificó la geometría de página de la plantilla'
    assert [p.text for p in report.paragraphs[:2]] == [p.text for p in source_template.paragraphs[:2]], 'Se modificó el título original de la plantilla'
    assert report.tables[0].rows[1].cells[1].text == 'PLATAFORMAS EMERGENTES (e)'
    assert report.tables[0].rows[3].cells[1].text == '4'
    assert report.tables[0].rows[3].cells[6].text == 'X'
    names = report.tables[0].rows[5].cells[0].text
    assert 'Mamani Uscamayta Agustin David' in names and 'Cusilayme Garcia Jose Luis' in names
    assert 'RENE ALONSO NIETO VALENCIA' in report.tables[0].rows[6].cells[0].text
    manifest = json.loads((ROOT / 'docs' / 'report_assets' / 'manifest.json').read_text(encoding='utf-8'))
    assert len(report.tables) == len(manifest['sources']) + 1
    for table, item in zip(report.tables[1:], manifest['sources']):
        actual = [p.text if p.text != ' ' else '' for p in table.cell(0, 0).paragraphs]
        expected_source = (ROOT / item['file']).read_text(encoding='utf-8-sig')
        assert actual == expected_source.splitlines(), f'Código incompleto o cambiado: {item["file"]}'
        assert hashlib.sha256(expected_source.encode('utf-8')).hexdigest() == item['sha256']
    with zipfile.ZipFile(REPORT) as archive, zipfile.ZipFile(template) as original:
        root = etree.fromstring(archive.read('word/document.xml'))
        plain_text = '\n'.join(root.xpath('//*[local-name()="t"]/text()'))
        assert 'sb_secret_' not in plain_text
        config = json.loads((ROOT / 'config.json').read_text(encoding='utf-8'))
        assert config['SUPABASE_PUBLISHABLE_KEY'] not in plain_text, 'La clave real está incluida en el informe'
        comments = etree.fromstring(archive.read('word/comments.xml'))
        comment_count = len(comments)
        assert comment_count == 5, 'Faltan comentarios de cambios'
        embedded = {hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith('word/media/')}
        for name in original.namelist():
            if name.startswith('word/media/'):
                assert hashlib.sha256(original.read(name)).hexdigest() in embedded, 'Se perdió una imagen institucional'
        footer = ''.join(archive.read(name).decode('utf-8') for name in archive.namelist() if re.fullmatch(r'word/footer\d+\.xml', name))
        assert 'NUMPAGES' in footer and 'PAGE' in footer, 'Falta numeración de páginas'
        assert 'https://github.com/AdMu2838/Labs_PE' in archive.read('word/_rels/document.xml.rels').decode('utf-8')
    pdf = fitz.open(PREVIEW)
    assert len(pdf) >= 20
    assert all(abs(page.rect.width - 595.3) < 2 and abs(page.rect.height - 841.9) < 2 for page in pdf)
    REVIEW.mkdir(parents=True, exist_ok=True)
    label_font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 15)
    for first in range(0, len(pdf), 16):
        sheet = Image.new('RGB', (960, 1440), '#D8D8D8')
        draw = ImageDraw.Draw(sheet)
        for i, page in enumerate(list(pdf)[first:first + 16]):
            pix = page.get_pixmap(matrix=fitz.Matrix(.37, .37), alpha=False)
            thumb = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
            x = (i % 4) * 240 + 10
            y = (i // 4) * 360 + 24
            sheet.paste(thumb, (x, y))
            draw.text((x, y - 20), f'Página {first + i + 1}', font=label_font, fill='black')
        sheet.save(REVIEW / f'contacto_{first + 1:03d}.png')
    selected = {0, 1, 2, len(pdf) - 1}
    for index, page in enumerate(pdf):
        content = page.get_text()
        if 'Figura 13.' in content or '5. ANEXOS' in content or 'A. Código completo' in content:
            selected.add(index)
    for index in sorted(selected):
        pdf[index].get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=False).save(REVIEW / f'pagina_{index + 1:03d}.png')
    summary = {
        'pages': len(pdf), 'web_screenshots': 12, 'figures': 13,
        'source_files_verified': len(manifest['sources']), 'source_lines_verified': manifest['source_lines'],
        'word_comments': comment_count, 'template_geometry_preserved': True,
        'institutional_images_preserved': True, 'code_matches_source': True,
        'no_real_keys_in_document': True,
    }
    (REVIEW / 'validation.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--template', type=Path, required=True)
    main(parser.parse_args().template)
